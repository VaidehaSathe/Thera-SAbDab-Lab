from __future__ import annotations
import argparse, json, shutil, sqlite3, sys
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]; sys.path.insert(0,str(ROOT))
from thera_curator.constants import CANONICAL_COLUMNS,COLUMN_TO_DB,DB_FIELDS
from thera_curator.db import SCHEMA_SQL
from thera_curator.utils import utcnow,sha256_file


def build(workbook:Path,out_db:Path,out_json:Path,pdf:Path|None=None):
    wb=load_workbook(workbook,read_only=True,data_only=False)
    ws=wb['Antibody Database']; headers=[ws.cell(1,c).value for c in range(1,20)]
    if headers!=CANONICAL_COLUMNS: raise SystemExit('Canonical schema mismatch')
    if out_db.exists(): out_db.unlink()
    con=sqlite3.connect(out_db); con.executescript(SCHEMA_SQL); now=utcnow(); ref=[]
    pdfhash=sha256_file(pdf) if pdf and pdf.exists() else ''
    for row in ws.iter_rows(min_row=2,max_col=19,values_only=True):
        if not row[0]: continue
        rec={COLUMN_TO_DB[h]:('' if v is None else v) for h,v in zip(CANONICAL_COLUMNS,row)}
        rec['heavy_chain_count']=int(rec['heavy_chain_count'] or 0); rec['light_chain_count']=int(rec['light_chain_count'] or 0); ref.append(rec.copy())
        vals=[rec.get(f,'') for f in DB_FIELDS]
        con.execute(f"insert into records({','.join(DB_FIELDS)},who_list,source_pdf_hash,source_pdf_name,source_provenance_json,created_at,updated_at) values({','.join('?' for _ in DB_FIELDS)},?,?,?,?,?,?)",vals+['135',pdfhash,pdf.name if pdf else 'pl135.pdf','{}',now,now])
    if 'Evidence Ledger' in wb.sheetnames:
        ev=wb['Evidence Ledger']; rows=ev.iter_rows(min_row=2,max_col=14,values_only=True)
        for r in rows:
            if not any(v is not None for v in r): continue
            con.execute('insert into evidence_ledger(drug_name,evidence_category,alias_searched,database_source,exact_query,result_summary,source_title,source_identifier,url,url_validation_status,access_date,evidence_type,confidence,notes) values(?,?,?,?,?,?,?,?,?,?,?,?,?,?)',tuple('' if v is None else str(v) for v in r))
    con.execute('insert or replace into app_meta(key,value) values(?,?)',('seed_source',workbook.name)); con.execute('insert or replace into app_meta(key,value) values(?,?)',('seed_created',now)); con.commit(); con.close()
    out_json.write_text(json.dumps(ref,ensure_ascii=False,indent=2),encoding='utf-8')
    return len(ref)

if __name__=='__main__':
    ap=argparse.ArgumentParser(); ap.add_argument('workbook',type=Path); ap.add_argument('--pdf',type=Path); ap.add_argument('--db',type=Path,default=ROOT/'resources'/'seed.sqlite'); ap.add_argument('--json',type=Path,default=ROOT/'resources'/'golden'/'pl135_reference.json'); a=ap.parse_args(); print('records',build(a.workbook,a.db,a.json,a.pdf))
