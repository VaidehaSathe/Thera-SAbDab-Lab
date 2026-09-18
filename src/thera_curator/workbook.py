from __future__ import annotations
from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from .constants import CANONICAL_COLUMNS, COLUMN_TO_DB, DB_TO_COLUMN
from .utils import utcnow


def inspect_workbook(path: str|Path) -> dict:
    wb=load_workbook(path, read_only=True, data_only=False)
    info={'sheets':wb.sheetnames}
    if 'Antibody Database' not in wb.sheetnames:
        raise ValueError("Workbook is missing required 'Antibody Database' sheet")
    ws=wb['Antibody Database']
    headers=[ws.cell(1,c).value for c in range(1,ws.max_column+1)]
    info['headers']=headers
    info['schema_ok']=headers[:len(CANONICAL_COLUMNS)]==CANONICAL_COLUMNS
    info['row_count']=max(0,ws.max_row-1)
    return info


def read_antibody_records(path: str|Path) -> list[dict]:
    wb=load_workbook(path, read_only=True, data_only=False)
    if 'Antibody Database' not in wb.sheetnames:
        raise ValueError("Workbook is missing required 'Antibody Database' sheet")
    ws=wb['Antibody Database']
    headers=[ws.cell(1,c).value for c in range(1,ws.max_column+1)]
    if headers[:len(CANONICAL_COLUMNS)] != CANONICAL_COLUMNS:
        raise ValueError('Workbook canonical schema mismatch; merge aborted')
    out=[]
    for row in ws.iter_rows(min_row=2, max_col=len(CANONICAL_COLUMNS), values_only=True):
        if not row[0]: continue
        rec={COLUMN_TO_DB[h]: ('' if v is None else v) for h,v in zip(CANONICAL_COLUMNS,row)}
        rec['heavy_chain_count']=int(rec['heavy_chain_count'] or 0)
        rec['light_chain_count']=int(rec['light_chain_count'] or 0)
        out.append(rec)
    return out


def export_database_workbook(db, path: str|Path):
    path=Path(path)
    wb=Workbook(); ws=wb.active; ws.title='Antibody Database'
    ws.append(CANONICAL_COLUMNS)
    for cell in ws[1]:
        cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='24536B'); cell.alignment=Alignment(vertical='center',wrap_text=True)
    for rec in db.list_records(limit=100000):
        ws.append([rec.get(COLUMN_TO_DB[h],'') for h in CANONICAL_COLUMNS])
    widths=[22,70,44,14,70,14,70,70,70,16,65,48,65,65,65,65,65,20,30]
    for i,w in enumerate(widths,1): ws.column_dimensions[ws.cell(1,i).column_letter].width=w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions

    qa=wb.create_sheet('QA Summary'); qa.append(['Metric','Value'])
    d=db.dashboard()
    for k,v in [('Exported UTC',utcnow()),('Records',d['records']),('Favorites',d['favorites']),('Pending Review',d['pending']),('QC Review',d['qc_review'])]: qa.append([k,v])
    for status,count in db.con.execute("select coalesce(nullif(url_validation_status,''),'Unspecified'),count(*) from evidence_ledger where trim(coalesce(url,''))<>'' group by coalesce(nullif(url_validation_status,''),'Unspecified') order by 1"):
        qa.append([f'Hyperlinks — {status}',count])
    ev=wb.create_sheet('Evidence Ledger'); ev_headers=['Drug name','Evidence category','Alias searched','Database/source','Exact query','Result summary','Source title','Source identifier','URL','URL validation status','Access date','Evidence type','Confidence','Notes']; ev.append(ev_headers)
    rows=db.con.execute('select drug_name,evidence_category,alias_searched,database_source,exact_query,result_summary,source_title,source_identifier,url,url_validation_status,access_date,evidence_type,confidence,notes from evidence_ledger order by id').fetchall()
    for r in rows: ev.append(list(r))
    ss=wb.create_sheet('Sources & Scope'); ss.append(['Item','Details']); ss.append(['Application','Thera-SAbDab WHO INN Curator']); ss.append(['Exported',utcnow()])
    mr=wb.create_sheet('Manual Review'); mr.append(['Drug name','Reason for manual review'])
    for r in db.con.execute("select drug_name,parse_status from pending_candidates where review_status='manual_review' or parse_status='REVIEW' order by drug_name"):
        mr.append([r[0],r[1]])
    for sh in wb.worksheets:
        sh.sheet_view.showGridLines=False
        for cell in sh[1]: cell.font=Font(bold=True)
    wb.save(path)
    return path
